"""
Excel Import/Export Utilities for Warehouse Stock Management
Replaces manual Excel tracking with direct database operations
"""

import sqlite3
from datetime import datetime, timedelta
from typing import List, Dict
import config


TEMPLATE_VERSION = "1.0"


SHEETS = {
    "INSTRUCTIONS": "Instructions",
    "LOCATIONS": "Locations",
    "SUPPLIERS": "Suppliers",
    "PRODUCTS": "Products",
    "STOCK_LEVELS": "StockLevels",
    "MOVEMENTS": "Movements",
}


PRODUCT_HEADERS = [
    "SKU",
    "Name",
    "Category",
    "Unit",
    "Unit Cost",
    "Unit Price",
    "Reorder Point",
    "Reorder Qty",
    "Max Stock",
    "Expiry Date (YYYY-MM-DD)",
    "Supplier Name",
    "Default Location (ZONE-AISLE-RACK-BIN)",
]

STOCK_HEADERS = [
    "SKU",
    "Location (ZONE-AISLE-RACK-BIN)",
    "Quantity",
    "Reserved",
    "Updated At (optional)",
]

LOCATION_HEADERS = [
    "Zone",
    "Aisle",
    "Rack",
    "Bin",
    "Capacity",
]

SUPPLIER_HEADERS = [
    "Name",
    "Contact Person",
    "Phone",
    "Email",
    "Lead Time Days",
    "Reliability Score (0-1)",
    "Active (TRUE/FALSE)",
]

MOVEMENT_HEADERS = [
    "Timestamp (YYYY-MM-DD HH:MM:SS or ISO)",
    "Type (IN/OUT/ADJUST/TRANSFER)",
    "SKU",
    "Location (ZONE-AISLE-RACK-BIN)",
    "To Location (ZONE-AISLE-RACK-BIN)",
    "Quantity",
    "Reference",
    "Notes",
    "User",
]


def _wb_styles():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    header_fill = PatternFill(start_color="2c3e50", end_color="2c3e50", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="bdc3c7")
    thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    vcenter = Alignment(vertical="center", wrap_text=True)

    return {
        "openpyxl": openpyxl,
        "header_fill": header_fill,
        "header_font": header_font,
        "thin_border": thin_border,
        "center": center,
        "vcenter": vcenter,
    }


def _apply_header(ws, headers):
    st = _wb_styles()
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = st["header_fill"]
        cell.font = st["header_font"]
        cell.alignment = st["center"]
        cell.border = st["thin_border"]

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(headers)).coordinate}"


def _autosize_columns(ws, max_width=50):
    for col_cells in ws.columns:
        max_len = 0
        col_letter = col_cells[0].column_letter
        for c in col_cells:
            try:
                if c.value is not None:
                    max_len = max(max_len, len(str(c.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, max_width)


def _norm_bool(v, default=True):
    if v is None:
        return default
    s = str(v).strip().lower()
    if s in ("1", "true", "yes", "y"):
        return True
    if s in ("0", "false", "no", "n"):
        return False
    return default


def _loc_label(zone, aisle, rack, bin_):
    return f"{zone}-{aisle}-{rack}-{bin_}"


def create_warehouse_stock_template(filepath: str):
    """
    Create a multi-sheet template workbook for manual warehouse stock tracking.
    This is the canonical format the app can import/export.
    """
    try:
        import openpyxl
    except ImportError:
        raise ImportError("openpyxl is required. Install with: pip install openpyxl")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Instructions
    ins = wb.create_sheet(SHEETS["INSTRUCTIONS"])
    ins["A1"] = "Warehouse Stock Sheet (Template)"
    ins["A2"] = f"Version: {TEMPLATE_VERSION}"
    ins["A4"] = "How to use:"
    ins["A5"] = "1) Fill Suppliers (optional) and Locations (required)."
    ins["A6"] = "2) Fill Products (SKU is required, unique)."
    ins["A7"] = "3) Fill StockLevels to set quantities per location."
    ins["A8"] = "4) Import this workbook in the app (Management → Import Excel)."
    ins["A10"] = "Notes:"
    ins["A11"] = "- Dates should be YYYY-MM-DD."
    ins["A12"] = "- Locations use the label: ZONE-AISLE-RACK-BIN (must match Locations sheet)."

    # Locations
    ws_loc = wb.create_sheet(SHEETS["LOCATIONS"])
    _apply_header(ws_loc, LOCATION_HEADERS)
    ws_loc.append(["MAIN", "01", "R1", "B1", 100000])

    # Suppliers
    ws_sup = wb.create_sheet(SHEETS["SUPPLIERS"])
    _apply_header(ws_sup, SUPPLIER_HEADERS)
    ws_sup.append(["Default Supplier", "", "", "", 7, 0.8, "TRUE"])

    # Products
    ws_prod = wb.create_sheet(SHEETS["PRODUCTS"])
    _apply_header(ws_prod, PRODUCT_HEADERS)

    # StockLevels
    ws_stock = wb.create_sheet(SHEETS["STOCK_LEVELS"])
    _apply_header(ws_stock, STOCK_HEADERS)

    # Movements (optional)
    ws_mov = wb.create_sheet(SHEETS["MOVEMENTS"])
    _apply_header(ws_mov, MOVEMENT_HEADERS)

    for ws in (ws_loc, ws_sup, ws_prod, ws_stock, ws_mov):
        _autosize_columns(ws)
    _autosize_columns(ins)

    wb.save(filepath)
    return filepath


def export_warehouse_stock_workbook(filepath: str, db_path=None):
    """Export database content into the canonical multi-sheet warehouse workbook."""
    try:
        import openpyxl
    except ImportError:
        raise ImportError("openpyxl is required. Install with: pip install openpyxl")

    db = db_path or config.DB_PATH
    conn = sqlite3.connect(db)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Instructions
    ins = wb.create_sheet(SHEETS["INSTRUCTIONS"])
    ins["A1"] = "Warehouse Stock Sheet (Export)"
    ins["A2"] = f"Version: {TEMPLATE_VERSION}"
    ins["A3"] = f"Exported at: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ins["A5"] = "This workbook can be re-imported into the app."

    # Locations
    ws_loc = wb.create_sheet(SHEETS["LOCATIONS"])
    _apply_header(ws_loc, LOCATION_HEADERS)
    cursor.execute("SELECT * FROM locations ORDER BY zone, aisle, rack, bin")
    for r in cursor.fetchall():
        ws_loc.append([r["zone"], r["aisle"], r["rack"], r["bin"], r["capacity"]])

    # Suppliers
    ws_sup = wb.create_sheet(SHEETS["SUPPLIERS"])
    _apply_header(ws_sup, SUPPLIER_HEADERS)
    cursor.execute("SELECT * FROM suppliers ORDER BY active DESC, name")
    for s in cursor.fetchall():
        ws_sup.append([
            s["name"],
            s["contact_person"] or "",
            s["phone"] or "",
            s["email"] or "",
            s["lead_time_days"],
            s["reliability_score"],
            "TRUE" if bool(s["active"]) else "FALSE",
        ])

    # Products
    ws_prod = wb.create_sheet(SHEETS["PRODUCTS"])
    _apply_header(ws_prod, PRODUCT_HEADERS)
    cursor.execute("""
        SELECT p.*, s.name as supplier_name, l.zone, l.aisle, l.rack, l.bin
        FROM products p
        LEFT JOIN suppliers s ON p.supplier_id = s.id
        LEFT JOIN locations l ON p.location_id = l.id
        ORDER BY p.category, p.name
    """)
    for p in cursor.fetchall():
        default_loc = _loc_label(p["zone"], p["aisle"], p["rack"], p["bin"]) if p["zone"] else ""
        ws_prod.append([
            p["sku"],
            p["name"],
            p["category"] or "",
            p["unit"] or "pcs",
            float(p["unit_cost"] or 0),
            float(p["unit_price"] or 0),
            int(p["reorder_point"] or 0),
            int(p["reorder_qty"] or 0),
            int(p["max_stock"] or 0),
            p["expiry_date"] or "",
            p["supplier_name"] or "",
            default_loc,
        ])

    # StockLevels
    ws_stock = wb.create_sheet(SHEETS["STOCK_LEVELS"])
    _apply_header(ws_stock, STOCK_HEADERS)
    cursor.execute("""
        SELECT sl.*, p.sku, l.zone, l.aisle, l.rack, l.bin
        FROM stock_levels sl
        JOIN products p ON sl.product_id = p.id
        JOIN locations l ON sl.location_id = l.id
        ORDER BY p.sku, l.zone, l.aisle, l.rack, l.bin
    """)
    for sl in cursor.fetchall():
        ws_stock.append([
            sl["sku"],
            _loc_label(sl["zone"], sl["aisle"], sl["rack"], sl["bin"]),
            int(sl["quantity"] or 0),
            int(sl["reserved"] or 0),
            sl["updated_at"] or "",
        ])

    # Movements (last 90 days)
    ws_mov = wb.create_sheet(SHEETS["MOVEMENTS"])
    _apply_header(ws_mov, MOVEMENT_HEADERS)
    cutoff = (datetime.now() - timedelta(days=90)).isoformat()
    cursor.execute("""
        SELECT sm.*, p.sku, l.zone, l.aisle, l.rack, l.bin,
               l2.zone as tz, l2.aisle as ta, l2.rack as tr, l2.bin as tb
        FROM stock_movements sm
        JOIN products p ON sm.product_id = p.id
        LEFT JOIN locations l ON sm.location_id = l.id
        LEFT JOIN locations l2 ON sm.to_location_id = l2.id
        WHERE sm.timestamp >= ?
        ORDER BY sm.timestamp DESC
        LIMIT 2000
    """, (cutoff,))
    for m in cursor.fetchall():
        loc = _loc_label(m["zone"], m["aisle"], m["rack"], m["bin"]) if m["zone"] else ""
        to_loc = _loc_label(m["tz"], m["ta"], m["tr"], m["tb"]) if m["tz"] else ""
        ws_mov.append([
            m["timestamp"],
            m["movement_type"],
            m["sku"],
            loc,
            to_loc,
            int(m["quantity"] or 0),
            m["reference"] or "",
            m["notes"] or "",
            m["user"] or "",
        ])

    for ws in (ws_loc, ws_sup, ws_prod, ws_stock, ws_mov):
        _autosize_columns(ws)
    _autosize_columns(ins)

    conn.close()
    wb.save(filepath)
    return filepath


def import_warehouse_stock_workbook(filepath: str, db_path=None, update_existing=True):
    """
    Import the canonical warehouse stock workbook (template/export format).
    - Upserts suppliers by name
    - Upserts locations by (zone,aisle,rack,bin)
    - Upserts products by SKU
    - Upserts stock_levels by (product_id, location_id)
    """
    try:
        import openpyxl
    except ImportError:
        raise ImportError("openpyxl is required. Install with: pip install openpyxl")

    db = db_path or config.DB_PATH
    conn = sqlite3.connect(db)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    wb = openpyxl.load_workbook(filepath, data_only=True)

    def sheet(name):
        return wb[name] if name in wb.sheetnames else None

    ws_loc = sheet(SHEETS["LOCATIONS"])
    ws_sup = sheet(SHEETS["SUPPLIERS"])
    ws_prod = sheet(SHEETS["PRODUCTS"])
    ws_stock = sheet(SHEETS["STOCK_LEVELS"])

    imported = {"suppliers": 0, "locations": 0, "products": 0, "stock_rows": 0}
    updated = {"suppliers": 0, "locations": 0, "products": 0, "stock_rows": 0}
    errors = []

    # --- Locations ---
    loc_id_by_label = {}
    if ws_loc:
        rows = list(ws_loc.iter_rows(min_row=2, values_only=True))
        for r in rows:
            try:
                zone, aisle, rack, bin_, cap = r[:5]
                if not zone or not aisle or not rack or not bin_:
                    continue
                cap = int(cap or 100)
                cursor.execute("""
                    SELECT id FROM locations
                    WHERE zone=? AND aisle=? AND rack=? AND bin=?
                """, (str(zone), str(aisle), str(rack), str(bin_)))
                ex = cursor.fetchone()
                if ex:
                    if update_existing:
                        cursor.execute("UPDATE locations SET capacity=? WHERE id=?", (cap, ex["id"]))
                        updated["locations"] += 1
                    loc_id = ex["id"]
                else:
                    cursor.execute("""
                        INSERT INTO locations(zone, aisle, rack, bin, capacity, current_load)
                        VALUES (?, ?, ?, ?, ?, 0)
                    """, (str(zone), str(aisle), str(rack), str(bin_), cap))
                    imported["locations"] += 1
                    loc_id = cursor.lastrowid
                loc_id_by_label[_loc_label(str(zone), str(aisle), str(rack), str(bin_))] = loc_id
            except Exception as e:
                errors.append(f"Locations row error: {e}")

    # If none provided, ensure at least one exists
    cursor.execute("SELECT id, zone, aisle, rack, bin FROM locations ORDER BY id ASC LIMIT 1")
    first_loc = cursor.fetchone()
    if first_loc:
        loc_id_by_label.setdefault(_loc_label(first_loc["zone"], first_loc["aisle"], first_loc["rack"], first_loc["bin"]), first_loc["id"])
    else:
        cursor.execute("INSERT INTO locations(zone, aisle, rack, bin, capacity, current_load) VALUES ('MAIN','01','R1','B1',100000,0)")
        loc_id = cursor.lastrowid
        imported["locations"] += 1
        loc_id_by_label[_loc_label("MAIN", "01", "R1", "B1")] = loc_id

    # --- Suppliers ---
    sup_id_by_name = {}
    if ws_sup:
        for r in ws_sup.iter_rows(min_row=2, values_only=True):
            try:
                name, contact, phone, email, lead, rel, active = (list(r) + [None] * 7)[:7]
                if not name:
                    continue
                name = str(name).strip()
                lead = int(lead or 7)
                rel = float(rel or 0.8)
                active_b = _norm_bool(active, default=True)
                cursor.execute("SELECT id FROM suppliers WHERE name = ?", (name,))
                ex = cursor.fetchone()
                if ex:
                    if update_existing:
                        cursor.execute("""
                            UPDATE suppliers
                            SET contact_person=?, phone=?, email=?, lead_time_days=?, reliability_score=?, active=?
                            WHERE id=?
                        """, (contact or "", phone or "", email or "", lead, rel, 1 if active_b else 0, ex["id"]))
                        updated["suppliers"] += 1
                    sup_id = ex["id"]
                else:
                    cursor.execute("""
                        INSERT INTO suppliers(name, contact_person, phone, email, lead_time_days, reliability_score, active)
                        VALUES (?, ?, ?, ?, ?, ?, ?)
                    """, (name, contact or "", phone or "", email or "", lead, rel, 1 if active_b else 0))
                    imported["suppliers"] += 1
                    sup_id = cursor.lastrowid
                sup_id_by_name[name] = sup_id
            except Exception as e:
                errors.append(f"Suppliers row error: {e}")

    # Ensure default supplier exists
    cursor.execute("SELECT id, name FROM suppliers WHERE active = TRUE ORDER BY id ASC LIMIT 1")
    first_sup = cursor.fetchone()
    if first_sup:
        sup_id_by_name.setdefault(first_sup["name"], first_sup["id"])
    else:
        cursor.execute("INSERT INTO suppliers(name, lead_time_days, reliability_score, active) VALUES ('Default Supplier', 7, 0.8, TRUE)")
        sup_id = cursor.lastrowid
        imported["suppliers"] += 1
        sup_id_by_name["Default Supplier"] = sup_id

    # --- Products ---
    prod_id_by_sku = {}
    if ws_prod:
        # Map headers by exact/contains
        headers = [c.value for c in ws_prod[1]]
        def idx(name):
            for i, h in enumerate(headers):
                if h and str(name).lower() in str(h).lower():
                    return i
            return None

        i_sku = idx("SKU")
        i_name = idx("Name")
        if i_sku is None or i_name is None:
            errors.append("Products sheet missing required headers (SKU, Name).")
        else:
            for r in ws_prod.iter_rows(min_row=2, values_only=True):
                try:
                    sku = r[i_sku]
                    name = r[i_name]
                    if not sku or not name:
                        continue
                    sku = str(sku).strip()
                    name = str(name).strip()

                    # pull optional columns by header
                    def get(col, default=None):
                        j = idx(col)
                        if j is None:
                            return default
                        v = r[j]
                        return default if v is None else v

                    category = str(get("Category", "") or "")
                    unit = str(get("Unit", "pcs") or "pcs")
                    unit_cost = float(get("Unit Cost", 0) or 0)
                    unit_price = float(get("Unit Price", 0) or 0)
                    reorder_point = int(get("Reorder Point", 20) or 20)
                    reorder_qty = int(get("Reorder Qty", 50) or 50)
                    max_stock = int(get("Max Stock", 200) or 200)
                    expiry_date = get("Expiry Date", None)
                    expiry_date = str(expiry_date).strip() if expiry_date else None
                    supplier_name = str(get("Supplier Name", "") or "").strip()
                    default_loc = str(get("Default Location", "") or "").strip()

                    supplier_id = sup_id_by_name.get(supplier_name) if supplier_name else None
                    location_id = loc_id_by_label.get(default_loc) if default_loc else None

                    cursor.execute("SELECT id FROM products WHERE sku = ?", (sku,))
                    ex = cursor.fetchone()
                    if ex:
                        prod_id = ex["id"]
                        if update_existing:
                            cursor.execute("""
                                UPDATE products SET
                                    name=?, category=?, unit=?, unit_cost=?, unit_price=?,
                                    reorder_point=?, reorder_qty=?, max_stock=?,
                                    supplier_id=?, location_id=?, expiry_date=?
                                WHERE id=?
                            """, (
                                name, category, unit, unit_cost, unit_price,
                                reorder_point, reorder_qty, max_stock,
                                supplier_id, location_id, expiry_date, prod_id
                            ))
                            updated["products"] += 1
                    else:
                        cursor.execute("""
                            INSERT INTO products
                            (sku, name, category, unit, unit_cost, unit_price, reorder_point, reorder_qty, max_stock, supplier_id, location_id, expiry_date, created_at)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """, (
                            sku, name, category, unit, unit_cost, unit_price,
                            reorder_point, reorder_qty, max_stock,
                            supplier_id, location_id, expiry_date, datetime.now().isoformat()
                        ))
                        imported["products"] += 1
                        prod_id = cursor.lastrowid

                    prod_id_by_sku[sku] = prod_id
                except Exception as e:
                    errors.append(f"Products row error: {e}")

    # --- Stock Levels ---
    if ws_stock:
        headers = [c.value for c in ws_stock[1]]
        def idxs(name):
            for i, h in enumerate(headers):
                if h and str(name).lower() in str(h).lower():
                    return i
            return None

        i_sku = idxs("SKU")
        i_loc = idxs("Location")
        i_qty = idxs("Quantity")
        i_res = idxs("Reserved")

        if i_sku is None or i_loc is None or i_qty is None:
            errors.append("StockLevels sheet missing required headers (SKU, Location, Quantity).")
        else:
            for r in ws_stock.iter_rows(min_row=2, values_only=True):
                try:
                    sku = r[i_sku]
                    loc_label = r[i_loc]
                    qty = r[i_qty]
                    reserved = r[i_res] if i_res is not None else 0
                    if not sku or not loc_label:
                        continue
                    sku = str(sku).strip()
                    loc_label = str(loc_label).strip()
                    qty = int(qty or 0)
                    reserved = int(reserved or 0)

                    product_id = prod_id_by_sku.get(sku)
                    if not product_id:
                        cursor.execute("SELECT id FROM products WHERE sku=?", (sku,))
                        ex = cursor.fetchone()
                        if not ex:
                            continue
                        product_id = ex["id"]
                        prod_id_by_sku[sku] = product_id

                    location_id = loc_id_by_label.get(loc_label)
                    if not location_id:
                        # try to create location from label
                        parts = loc_label.split("-")
                        if len(parts) == 4:
                            z, a, rk, b = [p.strip() for p in parts]
                            cursor.execute("""
                                INSERT OR IGNORE INTO locations(zone, aisle, rack, bin, capacity, current_load)
                                VALUES (?, ?, ?, ?, 100000, 0)
                            """, (z, a, rk, b))
                            cursor.execute("SELECT id FROM locations WHERE zone=? AND aisle=? AND rack=? AND bin=?", (z, a, rk, b))
                            loc_row = cursor.fetchone()
                            if loc_row:
                                location_id = loc_row["id"]
                                loc_id_by_label[loc_label] = location_id

                    if not location_id:
                        continue

                    cursor.execute("""
                        INSERT INTO stock_levels (product_id, location_id, quantity, reserved, updated_at)
                        VALUES (?, ?, ?, ?, ?)
                        ON CONFLICT(product_id, location_id) DO UPDATE SET
                            quantity = excluded.quantity,
                            reserved = excluded.reserved,
                            updated_at = excluded.updated_at
                    """, (product_id, location_id, qty, reserved, datetime.now().isoformat()))
                    imported["stock_rows"] += 1
                except Exception as e:
                    errors.append(f"StockLevels row error: {e}")

    conn.commit()
    conn.close()
    return {"imported": imported, "updated": updated, "errors": errors}


def export_products_to_excel(filepath: str, db_path=None):
    """
    Backwards compatible wrapper.
    Prefer `export_warehouse_stock_workbook()` which produces the canonical multi-sheet workbook.
    """
    return export_warehouse_stock_workbook(filepath, db_path=db_path)
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise ImportError("openpyxl is required. Install with: pip install openpyxl")
    
    db = db_path or config.DB_PATH
    conn = sqlite3.connect(db)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    
    # Get products with stock totals
    cursor.execute("""
        SELECT p.*, COALESCE(SUM(sl.quantity), 0) as total_stock,
               COALESCE(SUM(sl.reserved), 0) as total_reserved
        FROM products p
        LEFT JOIN stock_levels sl ON p.id = sl.product_id
        GROUP BY p.id
        ORDER BY p.category, p.name
    """)
    rows = [dict(row) for row in cursor.fetchall()]
    conn.close()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventory"
    
    # Header styling
    header_fill = PatternFill(start_color="2c3e50", end_color="2c3e50", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    headers = ["SKU", "Name", "Category", "Unit", "Unit Cost", "Unit Price", 
               "Reorder Point", "Reorder Qty", "Max Stock", "Current Stock", 
               "Reserved", "Available", "Status", "Expiry Date"]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    
    # Data rows
    for row_idx, row in enumerate(rows, 2):
        stock = row.get('total_stock', 0)
        reserved = row.get('total_reserved', 0)
        available = stock - reserved
        reorder_pt = row.get('reorder_point', 0)
        
        if stock <= 0:
            status = "STOCKOUT"
        elif stock <= reorder_pt:
            status = "BELOW REORDER"
        else:
            status = "OK"
        
        data = [
            row['sku'], row['name'], row.get('category', ''), row['unit'],
            row.get('unit_cost', 0), row.get('unit_price', 0),
            row.get('reorder_point', 0), row.get('reorder_qty', 0),
            row.get('max_stock', 0), stock, reserved, available,
            status, row.get('expiry_date', '')
        ]
        
        for col, val in enumerate(data, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
            
            # Color status
            if col == 13:  # Status column
                if val == "STOCKOUT":
                    cell.fill = PatternFill(start_color="e74c3c", end_color="e74c3c", fill_type="solid")
                    cell.font = Font(color="FFFFFF", bold=True)
                elif val == "BELOW REORDER":
                    cell.fill = PatternFill(start_color="f39c12", end_color="f39c12", fill_type="solid")
                    cell.font = Font(color="FFFFFF", bold=True)
                else:
                    cell.fill = PatternFill(start_color="27ae60", end_color="27ae60", fill_type="solid")
                    cell.font = Font(color="FFFFFF", bold=True)
    
    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        adjusted_width = min(max_length + 2, 40)
        ws.column_dimensions[column].width = adjusted_width
    
    # Freeze header
    ws.freeze_panes = "A2"
    
    wb.save(filepath)
    return len(rows)


def import_products_from_excel(filepath: str, db_path=None, update_existing=True):
    """
    Backwards compatible wrapper.
    If the workbook matches the canonical template, import everything.
    Otherwise, tries to interpret the active sheet as a simple products list.
    """
    try:
        import openpyxl
    except ImportError:
        raise ImportError("openpyxl is required. Install with: pip install openpyxl")

    wb = openpyxl.load_workbook(filepath, data_only=True)
    if SHEETS["PRODUCTS"] in wb.sheetnames or SHEETS["STOCK_LEVELS"] in wb.sheetnames:
        return import_warehouse_stock_workbook(filepath, db_path=db_path, update_existing=update_existing)

    # Fallback to legacy single-sheet import
    db = db_path or config.DB_PATH
    conn = sqlite3.connect(db)
    cursor = conn.cursor()
    ws = wb.active
    headers = [cell.value for cell in ws[1]]

    col_map = {}
    expected = ["SKU", "Name", "Category", "Unit", "Unit Cost", "Unit Price",
                "Reorder Point", "Reorder Qty", "Max Stock", "Current Stock", "Expiry Date"]

    for expected_col in expected:
        for idx, header in enumerate(headers, 1):
            if header and expected_col.lower() in str(header).lower():
                col_map[expected_col] = idx
                break

    imported = 0
    updated = 0
    errors = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        try:
            sku = row[col_map.get("SKU", 0) - 1] if "SKU" in col_map else None
            name = row[col_map.get("Name", 0) - 1] if "Name" in col_map else None
            if not sku or not name:
                continue

            category = row[col_map.get("Category", 0) - 1] if "Category" in col_map else ""
            unit = row[col_map.get("Unit", 0) - 1] if "Unit" in col_map else "pcs"
            unit_cost = row[col_map.get("Unit Cost", 0) - 1] if "Unit Cost" in col_map else 0.0
            unit_price = row[col_map.get("Unit Price", 0) - 1] if "Unit Price" in col_map else 0.0
            reorder_point = row[col_map.get("Reorder Point", 0) - 1] if "Reorder Point" in col_map else 20
            reorder_qty = row[col_map.get("Reorder Qty", 0) - 1] if "Reorder Qty" in col_map else 50
            max_stock = row[col_map.get("Max Stock", 0) - 1] if "Max Stock" in col_map else 200
            current_stock = row[col_map.get("Current Stock", 0) - 1] if "Current Stock" in col_map else 0
            expiry_date = row[col_map.get("Expiry Date", 0) - 1] if "Expiry Date" in col_map else None

            cursor.execute("SELECT id FROM products WHERE sku = ?", (str(sku),))
            existing = cursor.fetchone()

            if existing and update_existing:
                cursor.execute("""
                    UPDATE products SET
                        name = ?, category = ?, unit = ?, unit_cost = ?, unit_price = ?,
                        reorder_point = ?, reorder_qty = ?, max_stock = ?, expiry_date = ?
                    WHERE id = ?
                """, (str(name), str(category), str(unit), float(unit_cost or 0), float(unit_price or 0),
                      int(reorder_point or 0), int(reorder_qty or 0), int(max_stock or 0),
                      str(expiry_date) if expiry_date else None, existing[0]))
                updated += 1
            else:
                cursor.execute("""
                    INSERT INTO products (sku, name, category, unit, unit_cost, unit_price,
                        reorder_point, reorder_qty, max_stock, expiry_date, created_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (str(sku), str(name), str(category), str(unit), float(unit_cost or 0),
                      float(unit_price or 0), int(reorder_point or 0), int(reorder_qty or 0),
                      int(max_stock or 0), str(expiry_date) if expiry_date else None, datetime.now().isoformat()))
                imported += 1

                if current_stock and int(current_stock) > 0:
                    new_id = cursor.lastrowid
                    cursor.execute("""
                        INSERT OR IGNORE INTO locations(zone, aisle, rack, bin, capacity, current_load)
                        VALUES ('MAIN','01','R1','B1',100000,0)
                    """)
                    cursor.execute("SELECT id FROM locations ORDER BY id ASC LIMIT 1")
                    loc = cursor.fetchone()
                    loc_id = loc[0] if loc else 1
                    cursor.execute("""
                        INSERT INTO stock_levels (product_id, location_id, quantity, reserved, updated_at)
                        VALUES (?, ?, ?, 0, ?)
                        ON CONFLICT(product_id, location_id) DO UPDATE SET
                            quantity = excluded.quantity,
                            updated_at = excluded.updated_at
                    """, (new_id, loc_id, int(current_stock), datetime.now().isoformat()))

        except Exception as e:
            errors.append(f"Row error: {e}")

    conn.commit()
    conn.close()
    return {"imported": imported, "updated": updated, "errors": errors}


def export_stock_movements_to_excel(filepath: str, days: int = 30, db_path=None):
    """Export recent stock movements to Excel"""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise ImportError("openpyxl is required")
    
    db = db_path or config.DB_PATH
    conn = sqlite3.connect(db)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    
    from datetime import timedelta
    cutoff = (datetime.now() - timedelta(days=days)).isoformat()
    
    cursor.execute("""
        SELECT sm.*, p.sku, p.name, l.zone, l.aisle, l.rack, l.bin
        FROM stock_movements sm
        JOIN products p ON sm.product_id = p.id
        LEFT JOIN locations l ON sm.location_id = l.id
        WHERE sm.timestamp >= ?
        ORDER BY sm.timestamp DESC
    """, (cutoff,))
    
    rows = [dict(row) for row in cursor.fetchall()]
    conn.close()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock Movements"
    
    header_fill = PatternFill(start_color="2c3e50", end_color="2c3e50", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    headers = ["Timestamp", "Type", "SKU", "Product", "Quantity", "Location", "Reference", "Notes", "User"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
    
    type_colors = {
        'IN': '27ae60',
        'OUT': 'e74c3c',
        'ADJUST': 'f39c12',
        'TRANSFER': '9b59b6'
    }
    
    for row_idx, row in enumerate(rows, 2):
        loc = f"{row.get('zone','')}-{row.get('aisle','')}-{row.get('rack','')}-{row.get('bin','')}" if row.get('zone') else 'N/A'
        data = [
            row['timestamp'], row['movement_type'], row['sku'], row['name'],
            row['quantity'], loc, row.get('reference', ''), row.get('notes', ''), row.get('user', '')
        ]
        
        for col, val in enumerate(data, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = thin_border
            
            if col == 2:  # Type column
                color = type_colors.get(val, '95a5a6')
                cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
                cell.font = Font(color="FFFFFF", bold=True)
    
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[column].width = min(max_length + 2, 40)
    
    ws.freeze_panes = "A2"
    wb.save(filepath)
    return len(rows)
